"""Diapositiva de título de la presentación PPTX."""

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from functions.write_presentation.constants import (
    MARGIN_L,
    CONTENT_W,
)
from functions.write_presentation.helpers import (
    add_text_box,
    remove_placeholders,
)


def title_slide(presentacion: Presentation, institucion: str, titulacion: str, fecha: str):
    """
    Genera la diapositiva de título.

    Args:
        presentacion: Plantilla de PowerPoint.
        institucion: Nombre de la institución.
        titulacion: Nombre de la titulación.
        fecha: Fecha de la presentación.
    """
    layout = presentacion.slide_layouts[0]  # Layout de título
    slide = presentacion.slides.add_slide(layout)  # Añade la diapositiva
    remove_placeholders(slide)  # Elimina los marcadores de posición

    add_text_box(
        slide,
        titulacion,
        left=MARGIN_L,
        top=Inches(2.2),
        width=CONTENT_W,
        height=Inches(1.2),
        size=48,
        bold=True,
        color=(255, 255, 255),
        align=PP_ALIGN.CENTER,
    )  # Añade el texto de la titulación
    add_text_box(
        slide,
        institucion,
        left=MARGIN_L,
        top=Inches(3.6),
        width=CONTENT_W,
        height=Inches(0.6),
        size=36,
        color=(220, 220, 220),
        align=PP_ALIGN.CENTER,
    )  # Añade el texto de la institución
    add_text_box(
        slide,
        fecha,
        left=MARGIN_L,
        top=Inches(4.4),
        width=CONTENT_W,
        height=Inches(0.4),
        size=24,
        color=(180, 180, 180),
        align=PP_ALIGN.CENTER,
    )  # Añade el texto de la fecha
