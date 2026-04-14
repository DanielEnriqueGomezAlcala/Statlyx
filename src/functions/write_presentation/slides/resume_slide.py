"""Diapositiva de resumen de sección con gráfica y texto."""

from pptx import Presentation
from pptx.util import Inches

from functions.write_presentation.constants import (
    MARGIN_L,
    MARGIN_T,
    CONTENT_W,
    SLIDE_W,
)
from functions.write_presentation.helpers import (
    add_text_box,
    add_image,
    remove_placeholders,
)


def resume_slide(
    presentacion: Presentation, title: str, img_path: str | None, text: str | None
):
    """
    Genera la diapositiva de resumen de sección.

    Args:
        presentacion: Plantilla de PowerPoint.
        title: Título de la diapositiva.
        img_path: Ruta de la imagen.
        text: Texto de la diapositiva.
    """
    layout = presentacion.slide_layouts[1]  # Layout de resumen
    slide = presentacion.slides.add_slide(layout)
    remove_placeholders(slide)  # Elimina los marcadores de posición

    add_text_box(
        slide,
        title,
        left=MARGIN_L,
        top=MARGIN_T,
        width=CONTENT_W,
        height=Inches(0.55),
        size=18,
        bold=True,
    )  # Añade el texto del título

    top_img = Inches(1.15)
    img_w = Inches(9.0)
    img_left = (SLIDE_W - img_w) / 2

    if img_path:
        add_image(
            slide, img_path, img_left, top_img, img_w
        )  # Añadimos grafico de resumen

    if text:
        add_text_box(
            slide,
            text,
            left=MARGIN_L,
            top=Inches(5.0),
            width=CONTENT_W,
            height=Inches(2.3),
            size=12,
        )  # Añadimos texto del resumen
