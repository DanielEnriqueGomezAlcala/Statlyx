
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from functions.write_presentation.constants import (
    MARGIN_L, CONTENT_W,
)
from functions.write_presentation.helpers import (
    add_text_box, remove_placeholders,
)

def section_slide(presentacion: Presentation, title: str, subtitle: str = ""):
    layout = presentacion.slide_layouts[2]
    slide = presentacion.slides.add_slide(layout)
    remove_placeholders(slide)

    add_text_box(
        slide, title, left=MARGIN_L, top=Inches(2.8), width=CONTENT_W, height=Inches(1.0), size=36, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
    if subtitle:
        add_text_box(slide, subtitle, left=MARGIN_L, top=Inches(3.9), width=CONTENT_W, height=Inches(0.6), size=18, color=(220, 220, 220), align=PP_ALIGN.CENTER)