"""Diapositiva de conclusiones y recomendaciones."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from functions.write_presentation.constants import (
    MARGIN_L,
    MARGIN_T,
    CONTENT_W,
)
from functions.write_presentation.helpers import (
    add_text_box,
    remove_placeholders,
)

_GRAY = (120, 120, 120)
_DARK = (33, 33, 33)


def _add_label(slide, text: str, top):
    add_text_box(
        slide,
        text,
        left=MARGIN_L,
        top=top,
        width=CONTENT_W,
        height=Inches(0.32),
        size=10,
        bold=True,
        color=_GRAY,
    )


def _add_bullets(slide, bullets: list[str], top, height):
    text_box = slide.shapes.add_textbox(MARGIN_L, top, CONTENT_W, height)
    tf = text_box.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            paragraph = tf.paragraphs[0]
        else:
            paragraph = tf.add_paragraph()

        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = f"• {bullet}"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(*_DARK)
        paragraph.space_before = Pt(4)


def conclusions_slide(
    presentacion: Presentation,
    title: str,
    conclusion: str | None,
    bullets: list[str] | None,
):
    """
    Genera una diapositiva con conclusión en párrafo y recomendaciones en viñetas.

    Args:
        presentacion: Plantilla de PowerPoint.
        title: Título de la diapositiva.
        conclusion: Párrafo de conclusión.
        bullets: Lista de recomendaciones.
    """
    layout = presentacion.slide_layouts[1]
    slide = presentacion.slides.add_slide(layout)
    remove_placeholders(slide)

    # Título
    add_text_box(
        slide,
        title,
        left=MARGIN_L,
        top=MARGIN_T,
        width=CONTENT_W,
        height=Inches(0.55),
        size=18,
        bold=True,
    )

    current_top = Inches(1.15)

    # Bloque de conclusión
    if conclusion:
        _add_label(slide, "CONCLUSIÓN", current_top)
        current_top += Inches(0.38)
        add_text_box(
            slide,
            conclusion,
            left=MARGIN_L,
            top=current_top,
            width=CONTENT_W,
            height=Inches(2.2),
            size=12,
            color=_DARK,
        )
        current_top += Inches(2.3)

    # Bloque de recomendaciones
    if bullets:
        _add_label(slide, "RECOMENDACIONES", current_top)
        current_top += Inches(0.38)
        remaining = Inches(7.5) - current_top - Inches(0.2)
        _add_bullets(slide, bullets, current_top, remaining)
