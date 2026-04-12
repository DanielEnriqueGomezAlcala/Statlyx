"""Diapositiva de gráficas y texto de la presentación PPTX."""

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from functions.write_presentation.constants import (
    SLIDE_H,
    MARGIN_L,
    MARGIN_T,
    CONTENT_W,
)
from functions.write_presentation.helpers import (
    add_text_box,
    add_image,
    fit_image_in_area,
    remove_placeholders,
)


def chart_slide(
    presentation: Presentation,
    title: str,
    img_path: str | None,
    text: str | None,
    img_path2: str | None = None,
):
    """
    Genera la diapositiva de gráficos y texto.

    Args:
        presentation: Plantilla de PowerPoint.
        title: Título de la diapositiva.
        img_path: Ruta del grafico de linea.
        text: Texto de la diapositiva.
        img_path2: Ruta del grafico de tabla.
    """
    layout = presentation.slide_layouts[1]  # Layout de gráfico
    slide = presentation.slides.add_slide(layout)
    remove_placeholders(slide)  # Elimina los marcadores de posición

    add_text_box(
        slide,
        title,
        left=MARGIN_L,
        top=MARGIN_T,
        width=CONTENT_W,
        height=Inches(0.55),
        size=18,
        bold=True,
    )  # Añade el texto del título

    # Espacio disponible para el contenido (debajo del título, con margen inferior)
    content_top = Inches(1.15)
    available_height = SLIDE_H - content_top - Inches(0.22)

    text_reserved_height = Inches(1.1) if text else Inches(0)
    gap_between_chart_and_text = Inches(0.15) if text else Inches(0)
    gap_between_columns = Inches(0.2)

    chart_area_height = (
        available_height - text_reserved_height - gap_between_chart_and_text
    )

    if img_path and img_path2:
        # Dos imágenes
        right_col_width = int(CONTENT_W * 0.5)
        left_col_width = int(CONTENT_W - right_col_width - gap_between_columns)
        right_col_left = MARGIN_L + left_col_width + gap_between_columns

        # Ajustamos la imagen de la izquierda para que quepa en su columna sin desbordar
        left_img_left, left_img_top, left_img_width = fit_image_in_area(
            img_path, MARGIN_L, content_top, left_col_width, chart_area_height
        )
        add_image(slide, img_path, left_img_left, left_img_top, left_img_width)

        # Ajustamos la imagen de la derecha para que quepa en su columna sin desbordar
        right_img_left, right_img_top, right_img_width = fit_image_in_area(
            img_path2, right_col_left, content_top, right_col_width, chart_area_height
        )
        add_image(slide, img_path2, right_img_left, right_img_top, right_img_width)

        # Añadimos el texto si hay debajo de los graficos
        if text:
            text_top = content_top + chart_area_height + gap_between_chart_and_text
            text_height = SLIDE_H - text_top - Inches(0.22)
            add_text_box(
                slide,
                text,
                left=MARGIN_L,
                top=text_top,
                width=CONTENT_W,
                height=text_height,
                size=12,
            )

    elif img_path:
        # Un solo gráfico
        if text:  # Si hay texto a la columna derecha
            text_col_width = Inches(3.8)
            image_col_width = int(CONTENT_W - text_col_width - gap_between_columns)
            text_col_left = MARGIN_L + image_col_width + gap_between_columns
        else:
            # La imagen ocupa todo el ancho
            image_col_width = int(CONTENT_W)
            text_col_left = None

        # Ajustamos la imagen para que quepa en su columna sin desbordar
        img_left, img_top, img_width = fit_image_in_area(
            img_path, MARGIN_L, content_top, image_col_width, available_height
        )
        add_image(slide, img_path, img_left, img_top, img_width)

        # Añadimos el texto si hay a la derecha del gráfico
        if text and text_col_left is not None:
            add_text_box(
                slide,
                text,
                left=text_col_left,
                top=content_top,
                width=text_col_width,
                height=available_height,
                size=12,
                align=PP_ALIGN.JUSTIFY,
            )

    elif text:
        # El texto ocupa toda la diapositiva
        add_text_box(
            slide,
            text,
            left=MARGIN_L,
            top=content_top,
            width=CONTENT_W,
            height=available_height,
            size=14,
        )
