"""Diapositiva de información del análisis (período, tipologías y cursos)."""

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from functions.write_presentation.constants import (
    MARGIN_L,
    MARGIN_T,
    CONTENT_W,
)
from functions.write_presentation.helpers import (
    add_text_box,
    remove_placeholders,
)


def info_slide(presentacion: Presentation, rango_anios: str, tipologias: str, cursos: str):
    """
    Genera la diapositiva de información con datos del análisis realizado.

    Args:
        presentacion: Plantilla de PowerPoint.
        rango_anios: Rango de años analizados.
        tipologias: Tipologías analizadas.
        cursos: Cursos analizados.
    """
    layout = presentacion.slide_layouts[1]  # Layout de información
    slide = presentacion.slides.add_slide(layout)
    remove_placeholders(slide)  # Elimina los marcadores de posición

    add_text_box(
        slide,
        "Datos del análisis",
        left=MARGIN_L,
        top=MARGIN_T,
        width=CONTENT_W,
        height=Inches(0.6),
        size=24,
        bold=True,
    )  # Añade el texto del título

    lines = [  # Campos de información
        f"Período analizado:  {rango_anios}",
        f"Tipologías:  {tipologias}",
        f"Cursos:  {cursos}",
    ]

    add_text_box(
        slide,
        "\n".join(lines),
        left=MARGIN_L,
        top=Inches(1.3),
        width=CONTENT_W,
        height=Inches(2),
        size=24,
        align=PP_ALIGN.JUSTIFY,
    )  # Añade el texto de la información
