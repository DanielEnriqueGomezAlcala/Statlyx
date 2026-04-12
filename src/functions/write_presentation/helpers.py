import os

from PIL import Image as PILImage
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def set_text(
    text_frame, text: str, size: int, bold: bool, color: tuple, align: PP_ALIGN
):
    """
    Modifica el texto de cuadro de texto ya existente.

    Args:
        text_frame: Cuadro de texto.
        text: Texto a poner.
        size: Tamaño del texto.
        bold: Si es negrita o no.
        color: Color del texto.
        align: Alineación del texto (izquierda, centro, derecha).
    """
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    text_run = paragraph.add_run()
    text_run.text = text
    text_run.font.size = Pt(size)
    text_run.font.bold = bold
    text_run.font.color.rgb = RGBColor(*color)


def add_text_box(
    slide,
    text: str,
    left,
    top,
    width,
    height,
    size: int = 14,
    bold: bool = False,
    color: tuple = (33, 33, 33),
    align=PP_ALIGN.LEFT,
    wrap: bool = True,
):
    """
    Añade un cuadro de texto a la presentación.

    Args:
        slide: Diapositiva.
        text: Texto a poner.
        left: Posición izquierda del cuadro de texto.
        top: Posición superior del cuadro de texto.
        width: Ancho del cuadro de texto.
        height: Alto del cuadro de texto.
        size: Tamaño del texto.
        bold: Si es negrita o no.
        color: Color del texto.
        align: Alineación del texto (izquierda, centro, derecha).
        wrap: Si el texto se debe ajustar al ancho del cuadro de texto.
    Returns:
        Cuadro de texto por si es necesario modificarlo después.
    """
    text_box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = text_box.text_frame
    text_frame.word_wrap = wrap
    set_text(text_frame, text, size=size, bold=bold, color=color, align=align)
    return text_box


def add_image(slide, img_path: str, left, top, width):
    """
    Añade una imagen a la presentación.

    Args:
        slide: Diapositiva.
        img_path: Ruta de la imagen.
        left: Posición izquierda de la imagen.
        top: Posición superior de la imagen.
        width: Ancho de la imagen.
    """
    if img_path and os.path.exists(img_path):
        slide.shapes.add_picture(img_path, left, top, width=width)


def real_height(img_path: str, width) -> float:
    """
    Calcula el alto de una imagen para que se ajuste al ancho proporcionado.

    Args:
        img_path: Ruta de la imagen.
        width: Ancho de la imagen.
    Returns:
        Alto de la imagen.
    """
    if not img_path or not os.path.exists(img_path):
        return 0
    with PILImage.open(img_path) as image:
        original_width, original_height = image.size
    return int(width * original_height / original_width)


def fit_image_in_area(
    img_path: str, col_left, content_top, col_width, area_height
) -> tuple:
    """
    Calcula la posición y dimensiones de una imagen para que encaje centrada dentro de un área dada.

    Args:
        img_path: Ruta de la imagen.
        col_left: Posición izquierda de la columna donde se colocará la imagen.
        content_top: Posición superior del área de contenido.
        col_width: Ancho máximo disponible para la imagen.
        area_height: Alto máximo disponible para la imagen.
    Returns:
        Posición izquierda, posición superior y ancho de la imagen.
    """
    img_width = col_width
    img_height = real_height(img_path, img_width)

    if img_height > area_height:
        with PILImage.open(img_path) as image:
            original_width, original_height = image.size
        img_width = int(area_height * original_width / original_height)
        img_height = area_height

    img_left = col_left + (col_width - img_width) // 2
    img_top = content_top + (area_height - img_height) // 2
    return img_left, img_top, img_width


def remove_placeholders(slide):
    """
    Elimina los marcadores de posición de la diapositiva.

    Args:
        slide: Diapositiva.
    """
    for placeholder in list(slide.placeholders):
        placeholder_element = placeholder._element
        placeholder_element.getparent().remove(placeholder_element)
