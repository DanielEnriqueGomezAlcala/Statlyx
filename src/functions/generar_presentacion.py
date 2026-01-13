from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches as PptxInches
from dash import html
from functions.charts.interactive import (
    crear_lineas_rendimiento,
    crear_comparacion_media,
    crear_tabla_resumen,
    crear_brecha_genero
)


def generar_presentacion_pptx(carrera, universidad, df_filtrado):
    """
    Genera una presentación PPTX con los gráficos de rendimiento académico
    
    Args:
        carrera (str): Nombre de la carrera
        universidad (str): Nombre de la universidad
        df_filtrado (pd.DataFrame): DataFrame filtrado con los datos de la carrera
    
    Returns:
        tuple: (success, message, pptx_path o None)
    """
    try:
        # Crear directorios de salida
        base_path = Path(__file__).parent.parent.parent
        output_dir_graph = base_path / "output" / "graph"
        output_dir_graph.mkdir(parents=True, exist_ok=True)
        
        # Crear y guardar gráficos como imágenes
        fig1 = crear_lineas_rendimiento(df_filtrado, carrera)
        fig2 = crear_comparacion_media(df_filtrado, carrera)
        fig3 = crear_tabla_resumen(df_filtrado, carrera)
        fig4 = crear_brecha_genero(df_filtrado, carrera)
        
        # Guardar como imágenes PNG
        carrera_file = carrera.replace(' ', '_').replace('/', '_')
        img1_path = output_dir_graph / f"pres_rendimiento_{carrera_file}.png"
        img2_path = output_dir_graph / f"pres_comparacion_{carrera_file}.png"
        img3_path = output_dir_graph / f"pres_tabla_{carrera_file}.png"
        img4_path = output_dir_graph / f"pres_brecha_{carrera_file}.png"
        
        fig1.write_image(str(img1_path), width=1400, height=700)
        fig2.write_image(str(img2_path), width=1400, height=700)
        fig3.write_image(str(img3_path), width=1400, height=700)
        fig4.write_image(str(img4_path), width=1400, height=700)
        
        # Cargar plantilla PPTX
        template_path = base_path / "templates" / "Presentacion-calidad.pptx"
        prs = Presentation(str(template_path))

        wildcard_mappings = {
            "%carrera%": carrera,
            "%universidad%": universidad.replace('_', ' '),
            "%fecha%": datetime.now().strftime("%d/%m/%Y"),
            "%autor%": "Generado por Sistema de Análisis de Rendimiento Académico",
        }
        
        # Reemplazar wildcards y placeholders de imágenes en todas las diapositivas
        image_placeholders = {
            "LineasRendimiento": img1_path,
            "ComparacionMedia": img2_path,
            "TablaResumen": img3_path,
            "BrechaGenero": img4_path,
        }
        
        for slide in prs.slides:
            shapes_to_remove = []
            shapes_info = []
            
            for shape in slide.shapes:
                # Reemplazar wildcards en texto
                if hasattr(shape, "text_frame"):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            for wildcard, value in wildcard_mappings.items():
                                if wildcard in run.text:
                                    run.text = run.text.replace(wildcard, value)
                    
                    # Buscar placeholders de imágenes
                    shape_text = shape.text.strip()
                    if shape_text in image_placeholders:
                        # Guardar posición y tamaño del placeholder
                        shapes_info.append({
                            'image_path': image_placeholders[shape_text],
                            'left': shape.left,
                            'top': shape.top,
                            'width': shape.width,
                            'height': shape.height
                        })
                        shapes_to_remove.append(shape)
            
            # Eliminar placeholders de texto
            for shape in shapes_to_remove:
                sp = shape.element
                sp.getparent().remove(sp)
            
            # Agregar imágenes en las posiciones de los placeholders
            for info in shapes_info:
                slide.shapes.add_picture(
                    str(info['image_path']),
                    info['left'],
                    info['top'],
                    width=info['width'],
                    height=info['height']
                )
        
        # Guardar presentación
        output_dir_pptx = base_path / "output" / "presentacion_pptx"
        output_dir_pptx.mkdir(parents=True, exist_ok=True)
        pptx_path = output_dir_pptx / f"Presentacion_{carrera_file}.pptx"
        prs.save(str(pptx_path))
        
        return True, "Presentación PPTX generada exitosamente", pptx_path
        
    except Exception as e:
        import traceback
        error_detail = traceback.format_exc()
        print(f"Error detallado: {error_detail}")
        return False, str(e), None
